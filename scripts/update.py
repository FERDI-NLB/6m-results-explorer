"""Fail-closed official NLB result updater. Run by GitHub Actions every hour."""
import datetime as dt
import json
import re
import sys
from pathlib import Path

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / 'data.xlsx'
SITE = ROOT / 'docs' / 'index.html'
STATUS_FILE = ROOT / 'docs' / 'status.json'
URL = 'https://www.nlb.lk/results/mahajana-sampatha'
BASELINE = 6348
HEADERS = ['Date','Day','No','win','letter','1st digit','2nd digit','3rd digit','4th digit','5th digit','6th digit']
SL_TZ = dt.timezone(dt.timedelta(hours=5, minutes=30))


def now_iso():
    return dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat().replace('+00:00', 'Z')


def read_status():
    try:
        data = json.loads(STATUS_FILE.read_text(encoding='utf-8'))
        if isinstance(data, dict): return data
    except (OSError, ValueError):
        pass
    return {'state':'waiting','baseline_draw':BASELINE,'latest_verified_draw':BASELINE,
            'last_attempt_at':None,'last_successful_check_at':None,'last_updated_at':None,
            'last_updated_draws':[],'message':'Waiting for the first automated result check.'}


def save_status(status):
    STATUS_FILE.parent.mkdir(parents=True, exist_ok=True)
    temp = STATUS_FILE.with_suffix('.json.tmp')
    temp.write_text(json.dumps(status, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    temp.replace(STATUS_FILE)


def parse_official(html):
    """Only accept complete draw records explicitly grouped on the official page."""
    soup = BeautifulSoup(html, 'html.parser')
    found = {}
    for node in soup.select('article, .result, .result-item, .draw-result, .lottery-result, tr, li, section'):
        t = node.get_text(' ', strip=True)
        if len(t) > 500 or ('mahajana' not in t.lower() and 'draw' not in t.lower()):
            continue
        number = re.search(r'(?i)(?:draw\s*(?:no\.?|number|#)\s*[:#-]?\s*)(\d{4,5})', t)
        date = re.search(r'\b(20\d\d[-/]\d\d[-/]\d\d)\b', t)
        combo = re.search(r'(?i)(?:result|winning\s*(?:numbers|combination))\s*[:\-]?\s*([A-Z])\s*[-\s]*([0-9](?:[\s-]*[0-9]){5})\b', t)
        if not (number and date and combo): continue
        num = int(number.group(1)); day = dt.date.fromisoformat(date.group(1).replace('/', '-'))
        digits = re.sub(r'\D', '', combo.group(2))
        if len(digits) != 6: continue
        value = {'number':num,'date':day.isoformat(),'letter':combo.group(1).upper(),'digits':digits}
        if num in found and found[num] != value:
            raise ValueError(f'Conflicting official entries for draw {num}')
        found[num] = value
    if not found:
        raise ValueError('Could not verify complete draw records in official NLB HTML. Parser needs validation against live markup; no data changes made.')
    return sorted(found.values(), key=lambda x:x['number'])


def validate(record, existing):
    n = record['number']; d = dt.date.fromisoformat(record['date'])
    if n in existing:
        if existing[n] != (record['date'], record['letter'], record['digits']):
            raise ValueError(f'Draw {n} conflicts with saved workbook; manual review required')
        return False
    if n <= BASELINE: return False
    if n > max(existing) + 1: raise ValueError(f'Missing intermediate draw(s) before {n}; stopping to avoid gaps')
    if d > dt.datetime.now(SL_TZ).date(): raise ValueError(f'Future draw date {d} for {n}')
    if not re.fullmatch('[A-Z]', record['letter']) or not re.fullmatch('[0-9]{6}', record['digits']):
        raise ValueError(f'Incomplete draw {n}')
    return True


def rebuild_html():
    wb = load_workbook(BOOK, read_only=True, data_only=True); ws = wb['6 M']
    rows=[]
    for cells in list(ws.values)[1:]:
        date=cells[0]
        if hasattr(date, 'date'): date=date.date()
        if hasattr(date, 'isoformat'): date=date.isoformat()
        vals=[date,*cells[1:11]]
        rows.append(['' if x is None else str(x) for x in vals])
    rows.sort(key=lambda r:int(r[2]) if r[2].isdigit() else -1, reverse=True)
    html=SITE.read_text(encoding='utf-8')
    pattern=r'(const headers=)(.*?)(,allRows=)(.*?)(;\s*const digitIndices=)'
    replacement=lambda m:m.group(1)+json.dumps(HEADERS,ensure_ascii=False)+m.group(3)+json.dumps(rows,ensure_ascii=False,separators=(',',':'))+m.group(5)
    html,count=re.subn(pattern,replacement,html,count=1,flags=re.S)
    if count != 1: raise ValueError('Cannot find embedded dataset in website; deployment blocked')
    SITE.write_text(html, encoding='utf-8')
    return len(rows)


def main():
    status = read_status()
    attempt = now_iso()
    status.update({'state':'checking','baseline_draw':BASELINE,'last_attempt_at':attempt,
                   'message':'Checking the official NLB results page.','last_error':None})
    save_status(status)
    try:
        wb = load_workbook(BOOK); ws=wb['6 M']; existing={}
        # Read formula results from Excel's cached values when comparing existing draws.
        value_wb = load_workbook(BOOK, data_only=True, read_only=True)
        value_ws = value_wb['6 M']
        for row in value_ws.iter_rows(min_row=2, values_only=True):
            if row[2] is None: continue
            n=int(row[2]); date=row[0]
            if hasattr(date, 'date'): date=date.date()
            day=date.isoformat() if hasattr(date, 'isoformat') else str(date)[:10]
            # Use the saved winning-number text, not digit-column formulas that may have stale cached errors.
            win = str(row[3]).strip().upper() if row[3] is not None else ''
            if re.fullmatch(r'[A-Z][0-9]{6}', win):
                existing[n]=(day, win[0], win[1:])
        value_wb.close()
        if BASELINE not in existing: raise ValueError('Baseline draw 6348 is missing from Excel workbook')
        print('Baseline 6348 (Excel):', existing[BASELINE])
        response=requests.get(URL, timeout=30, headers={'User-Agent':'Mozilla/5.0 (compatible; 6MResultsUpdater/1.0)'})
        response.raise_for_status()
        records=parse_official(response.text)
        updates=[]
        for rec in records:
            if validate(rec, existing):
                updates.append(rec)
                existing[rec['number']] = (rec['date'], rec['letter'], rec['digits'])
        success_time=now_iso()
        status['last_successful_check_at']=success_time
        status['latest_verified_draw']=max(existing) if existing else BASELINE
        if not updates:
            status.update({'state':'ok','message':'Official NLB page checked and parsed successfully. No new verified draw was available.'})
            save_status(status)
            print('No verified new results. Website and workbook unchanged.')
            return
        row_lookup={int(ws.cell(r,3).value):r for r in range(2,ws.max_row+1) if ws.cell(r,3).value is not None}
        for rec in updates:
            n=rec['number']; d=dt.date.fromisoformat(rec['date']); digits=rec['digits']
            row=row_lookup.get(n, ws.max_row+1)
            vals=[dt.datetime.combine(d,dt.time()), d.strftime('%A'), n, rec['letter']+digits, rec['letter'], *map(int,digits)]
            for col,val in enumerate(vals,1): ws.cell(row,col).value=val
            row_lookup[n]=row
            print('Added verified official result:', n, rec['letter']+digits)
        wb.save(BOOK)
        row_count=rebuild_html()
        updated_time=now_iso()
        status.update({'state':'updated','last_successful_check_at':success_time,'last_updated_at':updated_time,
                       'latest_verified_draw':max(existing),'last_updated_draws':[r['number'] for r in updates],
                       'message':f'Added {len(updates)} verified draw(s) and rebuilt the dataset ({row_count} rows).'})
        save_status(status)
        print(f'Updated {len(updates)} draw(s) and rebuilt website ({row_count} rows).')
    except Exception as exc:
        status=read_status()
        status.update({'state':'error','last_attempt_at':attempt,'last_error':str(exc),
                       'message':f'Update blocked: {exc}'})
        save_status(status)
        print(f'UPDATE BLOCKED: {exc}', file=sys.stderr)
        raise


if __name__ == '__main__':
    main()
